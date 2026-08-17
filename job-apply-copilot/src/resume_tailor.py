"""Tailor your master resume + write a cover letter for one specific job.

Ground rules baked into the prompt:
  - Never invent experience, employers, titles, or skills you didn't list
    in profile.yaml / your master resume. Tailoring = re-emphasizing and
    re-ordering what's true, not fabricating a match.
  - Not keyword-stuffed: it should read like a person who actually has
    these skills wrote it for this role, not like a bot echoing the job
    description back.
  - Humanized tone: plain, direct, first-person where appropriate, no
    corporate-AI filler ("In today's fast-paced environment...", "I am
    thrilled to apply...", excessive em-dashes, "leverage synergies").
    Short sentences over long ones. Specifics over generic praise.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path

import anthropic
from docx import Document
from pypdf import PdfReader

GENERATED_DIR = Path(__file__).resolve().parent.parent / "data" / "generated"

SYSTEM_PROMPT = """You help tailor a real person's resume and write a cover \
letter for a specific job posting. You must never invent employers, titles, \
dates, degrees, or skills that are not present in the master resume or the \
candidate's stated skills list - only re-emphasize, re-order, and rephrase \
what is genuinely there, and select which true skills/experiences to \
foreground based on relevance to this specific job (not every skill needs \
to be mentioned every time).

Write in a plain, humanized voice: short and medium sentences, concrete \
detail, no corporate-AI filler phrases (avoid things like "in today's \
fast-paced environment", "I am thrilled/excited to apply", "leverage \
synergies", "passionate about", excessive em-dashes, or generic praise of \
the company with no specifics). It should read like a real engineer wrote \
it, not like an ad.

Output in exactly this format, nothing else:

===RESUME===
<tailored resume, plain text, ready to format>

===COVER_LETTER===
<3-4 short paragraphs, first person, no "Dear Hiring Manager" boilerplate \
if the job posting names a team or hiring manager - otherwise a simple \
neutral opener is fine>
"""


@dataclass
class TailoredApplication:
    resume_text: str
    cover_letter_text: str


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if path.suffix.lower() == ".docx":
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    return path.read_text()


def tailor(
    master_resume_path: Path,
    skills_text: str,
    job_title: str,
    company: str,
    job_description: str,
    model: str = "claude-sonnet-5",
) -> TailoredApplication:
    master_resume_text = extract_text(master_resume_path)

    client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
    user_prompt = f"""MASTER RESUME:
{master_resume_text}

CANDIDATE'S STATED SKILLS (free text, not exhaustive - use judgment on relevance):
{skills_text}

JOB TITLE: {job_title}
COMPANY: {company}
JOB DESCRIPTION:
{job_description}
"""

    response = client.messages.create(
        model=model,
        max_tokens=4000,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )
    text = "".join(block.text for block in response.content if hasattr(block, "text"))

    resume_match = re.search(r"===RESUME===\s*(.*?)\s*===COVER_LETTER===", text, re.S)
    cover_match = re.search(r"===COVER_LETTER===\s*(.*)", text, re.S)
    if not resume_match or not cover_match:
        raise ValueError(f"Unexpected tailoring output format:\n{text[:500]}")

    return TailoredApplication(
        resume_text=resume_match.group(1).strip(),
        cover_letter_text=cover_match.group(1).strip(),
    )


def save_as_docx(text: str, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(out_path))
    return out_path


def tailor_and_save(
    job_id: int,
    master_resume_path: Path,
    skills_text: str,
    job_title: str,
    company: str,
    job_description: str,
) -> tuple[Path, str]:
    """Returns (tailored_resume_docx_path, cover_letter_text)."""
    result = tailor(master_resume_path, skills_text, job_title, company, job_description)
    resume_path = GENERATED_DIR / f"job_{job_id}_resume.docx"
    save_as_docx(result.resume_text, resume_path)
    return resume_path, result.cover_letter_text
